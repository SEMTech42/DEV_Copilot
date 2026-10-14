"""Logique de reference du flux Expense-Policy-Gate, lue depuis expense_policy.xlsx.
Verifie la feuille Tests. Usage: python3 gate_check.py
"""
from pathlib import Path
import openpyxl

wb = openpyxl.load_workbook(Path(__file__).resolve().parent / "expense_policy.xlsx")
policy = {r[0]: r for r in wb["Policy"].iter_rows(min_row=2, values_only=True)}


def gate(category, amount, has_receipt):
    p = policy.get(category)
    if p is None:
        return "deny", "Categorie inconnue"
    _, max_amount, receipt_above, allowed, _c = p
    if allowed != "Oui":
        return "deny", "Categorie non remboursable"
    if amount > max_amount:
        return "deny", f"Plafond depasse ({max_amount} EUR max)"
    if amount > receipt_above and has_receipt != "Oui":
        return "deny", "Justificatif obligatoire"
    return "allow", "Conforme"


bad = 0
for t in wb["Tests"].iter_rows(min_row=2, values_only=True):
    tid, prompt, cat, amt, rec, expected, _ = t
    got, why = gate(cat, amt, rec)
    ok = got == expected
    bad += not ok
    print(f"{'PASS' if ok else 'FAIL'} {tid:4} {prompt:40} -> {got:5} {why}")
raise SystemExit(1 if bad else 0)
