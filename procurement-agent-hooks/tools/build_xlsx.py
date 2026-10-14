"""Builds data/procurement_data.xlsx (one Excel *table* per CSV) for the Excel Online (Business) connector.
Run: python3 tools/build_xlsx.py
"""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

DATA = Path(__file__).resolve().parent.parent / "data"
SHEETS = {  # sheet/table name -> (csv, numeric columns, bool columns)
    "Vendors": ("vendors.csv", {"risk_score"}, set()),
    "Employees": ("employees.csv", {"approval_limit_eur"}, {"active"}),
    "Budgets": ("budgets.csv", {"fiscal_year", "budget_eur", "committed_eur", "remaining_eur"}, set()),
    "PurchaseOrders": ("purchase_orders_seed.csv", {"amount_eur"}, set()),
    "PolicyRules": ("policy_rules.csv", {"priority"}, set()),
}
AUDIT_COLS = ["timestamp", "event", "conversation_id", "user", "tool", "decision", "rule_id", "reason", "input_json"]

wb = Workbook()
wb.remove(wb.active)


def add_table(name, header, rows):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r)
    last = ws.cell(row=ws.max_row, column=len(header)).coordinate
    t = Table(displayName=name, ref=f"A1:{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for i, h in enumerate(header, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(h) + 4)


for name, (fname, nums, bools) in SHEETS.items():
    with open(DATA / fname, newline="") as f:
        rd = csv.DictReader(f)
        header = rd.fieldnames
        rows = []
        for r in rd:
            row = []
            for h in header:
                v = r[h]
                if h in nums:
                    v = float(v) if "." in v else int(v)
                elif h in bools:
                    v = v == "true"
                row.append(v)
            rows.append(row)
    add_table(name, header, rows)
add_table("AuditLog", AUDIT_COLS, [["2026-10-14T00:00:00Z", "Init", "-", "-", "-", "log", "-", "Seed row (a table needs one data row); delete after the first real entry", "{}"]])
out = DATA / "procurement_data.xlsx"
wb.save(out)
print("written", out)
